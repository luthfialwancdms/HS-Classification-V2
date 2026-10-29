"""Debug: find why bab 27, 72, 84.71, 84.29 are missing"""
from docx import Document
import re

doc = Document(r'd:\Work\Project\Project_AI\docs\26~PMK.010~2022 BTKI 2022.docx')

# Search for specific codes in ALL tables
target_codes = ['8471', '8429', '8711', '2709', '2710', '7208', '7209']

print('=== SEARCHING ALL TABLES FOR TARGET CODES ===')
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows[:3]):  # check first 3 rows
        row_text = ' '.join(c.text.strip() for c in row.cells)
        for code in target_codes:
            if code in row_text:
                print(f'\nTable {t_idx}, Row {r_idx}: Found {code}')
                print(f'  Cols: {len(table.columns)}, Rows: {len(table.rows)}')
                # print first header row
                print(f'  Header: {[c.text.strip()[:25] for c in table.rows[0].cells]}')
                # print this row
                print(f'  Row {r_idx}: {[c.text.strip()[:35] for c in row.cells]}')
                break

print()
# Check what table structure bab 27 uses
print('=== LOOKING FOR BAB 27 (2709, 2710) ===')
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            txt = cell.text.strip()
            if txt.startswith('2709') or txt.startswith('2710') or txt == '27.09' or txt == '27.10':
                print(f'Table {t_idx}, Row {r_idx}, Cell: "{txt}"')
                print(f'  Full row: {[c.text.strip()[:30] for c in row.cells]}')
                print(f'  Table cols: {len(table.columns)}, rows: {len(table.rows)}')
                print(f'  Header row 0: {[c.text.strip()[:25] for c in table.rows[0].cells]}')
                break
