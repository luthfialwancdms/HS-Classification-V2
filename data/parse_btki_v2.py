"""
BTKI 2022 Parser v2
Handles multi-entry cells (multiple codes in one cell separated by newlines)
"""
import json, csv, re
from pathlib import Path
from docx import Document
from collections import Counter, defaultdict

DOCX_PATH = Path(r'd:\Work\Project\Project_AI\docs\26~PMK.010~2022 BTKI 2022.docx')
OUT_JSON    = Path(r'd:\Work\Project\Project_AI\data\btki_2022_full.json')
OUT_CSV     = Path(r'd:\Work\Project\Project_AI\data\btki_2022_full.csv')
OUT_NATIONAL= Path(r'd:\Work\Project\Project_AI\hs-classifier\src\data\btki_national.json')
OUT_REPORT  = Path(r'd:\Work\Project\Project_AI\data\btki_parse_report.md')

HS8 = re.compile(r'\d{4}\.\d{2}\.\d{2}')
HS6 = re.compile(r'\d{4}\.\d{2}')
HS4 = re.compile(r'\d{2}\.\d{2}')

def clean(t):
    t = t.replace('\t', ' ').strip()
    t = re.sub(r'[\u25a0\u25cf\ufffd]', '', t)   # remove OCR boxes
    t = re.sub(r' {2,}', ' ', t)
    return t

def extract_codes_from_cell(cell_text):
    """Extract all HS codes from a cell (may contain multiple)."""
    codes = []
    for line in cell_text.split('\n'):
        line = line.strip()
        # try 8-digit first
        m = HS8.search(line)
        if m:
            codes.append(('8', m.group()))
            continue
        m = HS6.search(line)
        if m:
            codes.append(('6', m.group()))
            continue
        m = HS4.search(line)
        if m:
            codes.append(('4', m.group()))
            continue
    return codes

def extract_texts_from_cell(cell_text):
    """Split cell text by newlines, return non-empty lines."""
    return [clean(l) for l in cell_text.split('\n') if clean(l)]

def extract_bm_from_cell(cell_text):
    """Extract bea masuk values, one per line matching codes."""
    lines = [clean(l) for l in cell_text.split('\n')]
    results = []
    for l in lines:
        if re.search(r'\d[\d,\.]*\s*%', l):
            results.append(l.strip())
        else:
            results.append('')
    return results

def get_seksi(bab):
    mapping = [(range(1,6),'I'),(range(6,15),'II'),(range(15,16),'III'),
               (range(16,25),'IV'),(range(25,28),'V'),(range(28,39),'VI'),
               (range(39,41),'VII'),(range(41,44),'VIII'),(range(44,47),'IX'),
               (range(47,50),'X'),(range(50,64),'XI'),(range(64,68),'XII'),
               (range(68,71),'XIII'),(range(71,72),'XIV'),(range(72,84),'XV'),
               (range(84,86),'XVI'),(range(86,90),'XVII'),(range(90,93),'XVIII'),
               (range(93,94),'XIX'),(range(94,97),'XX'),(range(97,100),'XXI')]
    for r,s in mapping:
        if bab in r: return s
    return '?'

def node_type(code):
    d = code.replace('.','')
    if len(d)==8: return 'tariff_line'
    if len(d)==6: return 'subheading'
    if len(d)==4: return 'heading'
    return 'other'

def is_tariff_table(table):
    if len(table.columns) < 4 or len(table.rows) < 3:
        return False
    header = ' '.join(c.text for c in table.rows[0].cells).lower()
    return 'pos' in header or 'hs code' in header or 'uraian' in header

def parse_table(table, seen, errors):
    """Parse one tariff table using zip-by-line approach for multi-entry cells."""
    records = []

    for r_idx, row in enumerate(table.rows):
        if r_idx < 2:
            continue  # skip header rows

        cells = [c.text for c in row.cells]
        if len(cells) < 4:
            continue

        hs_cell  = cells[1] if len(cells) > 1 else ''
        id_cell  = cells[2] if len(cells) > 2 else ''
        en_cell  = cells[3] if len(cells) > 3 else ''
        bm_cell  = cells[4] if len(cells) > 4 else ''

        hs_codes = extract_codes_from_cell(hs_cell)
        id_lines = extract_texts_from_cell(id_cell)
        en_lines = extract_texts_from_cell(en_cell)
        bm_lines = extract_bm_from_cell(bm_cell)

        if not hs_codes:
            continue

        # zip codes with matching description lines
        for i, (digits, code) in enumerate(hs_codes):
            # pick corresponding description line
            uraian_id = id_lines[i] if i < len(id_lines) else (id_lines[0] if id_lines else '')
            uraian_en = en_lines[i] if i < len(en_lines) else (en_lines[0] if en_lines else '')
            bm = bm_lines[i] if i < len(bm_lines) else ''
            # if bm empty, try finding first non-empty
            if not bm:
                for b in bm_lines:
                    if b: bm = b; break

            canon = code.replace('.','')
            bab_str = canon[:2]
            try:
                seksi = get_seksi(int(bab_str))
            except:
                seksi = '?'

            record = {
                'formatted_code': code,
                'canonical_code': canon,
                'digits': len(canon),
                'node_type': node_type(code),
                'bab': bab_str,
                'seksi': seksi,
                'uraian_id': clean(uraian_id),
                'uraian_en': clean(uraian_en),
                'bea_masuk_mfn': bm,
            }

            if canon not in seen:
                seen[canon] = record
                records.append(record)
            # else: keep first occurrence (usually cleaner)

    return records

def main():
    print('Loading DOCX...')
    doc = Document(str(DOCX_PATH))
    print(f'Tables: {len(doc.tables)}')

    seen = {}   # canonical_code -> record (dedup)
    errors = []
    table_count = 0

    for t_idx, table in enumerate(doc.tables):
        if is_tariff_table(table):
            parse_table(table, seen, errors)
            table_count += 1
        if t_idx % 50 == 0:
            print(f'  ... table {t_idx}/{len(doc.tables)}, codes so far: {len(seen)}')

    all_records = sorted(seen.values(), key=lambda x: x['canonical_code'])
    print(f'\nParsed {table_count} tariff tables → {len(all_records)} unique codes')

    lines8 = [r for r in all_records if r['digits'] == 8]
    lines6 = [r for r in all_records if r['digits'] == 6]
    lines4 = [r for r in all_records if r['digits'] == 4]
    print(f'  8-digit: {len(lines8)}, 6-digit: {len(lines6)}, 4-digit: {len(lines4)}')

    # save full JSON
    with open(OUT_JSON, 'w', encoding='utf-8') as f:
        json.dump({'note':'BTKI 2022 parsed from PMK 26/2022 DOCX',
                   'total': len(all_records), 'generated_at':'2026-08-11',
                   'data': all_records}, f, ensure_ascii=False, indent=2)
    print(f'Saved: {OUT_JSON}')

    # save CSV
    fields = ['formatted_code','canonical_code','digits','node_type','bab','seksi',
              'uraian_id','uraian_en','bea_masuk_mfn']
    with open(OUT_CSV, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(all_records)
    print(f'Saved: {OUT_CSV}')

    # update btki_national.json with real data
    national = [{'hs8': r['formatted_code'], 'uraian': r['uraian_id']}
                for r in lines8 if r['uraian_id']]
    with open(OUT_NATIONAL, 'w', encoding='utf-8') as f:
        json.dump({'note':'Subpos nasional BTKI 2022 — parsed dari DOCX PMK 26/2022',
                   'total_entries': len(national), 'generated_at':'2026-08-11',
                   'data': national}, f, ensure_ascii=False, indent=2)
    print(f'Updated btki_national.json: {len(national)} tariff lines')

    # verify key codes
    by_code = {r['formatted_code']: r for r in lines8}
    test = ['8471.30.20','8471.30.10','8704.10.31','8905.90.90',
            '8429.52.00','8711.20.90','2709.00.00','7208.51.00']
    print('\n=== KEY CODE VERIFICATION ===')
    for code in test:
        if code in by_code:
            print(f'  OK  {code} | {by_code[code]["uraian_id"][:60]}')
        else:
            print(f'  --- {code} MISSING')

    # bab coverage
    bab_counts = Counter(r['bab'] for r in lines8)
    print(f'\nBab coverage: {len(bab_counts)} chapters with tariff lines')
    missing_bab = [str(i).zfill(2) for i in range(1,98) if str(i).zfill(2) not in bab_counts]
    print(f'Missing bab: {missing_bab}')

    # write report
    report = f"""# BTKI 2022 Parse Report v2
Generated: 2026-08-11

## Summary
| Item | Count |
|------|-------|
| Tariff tables parsed | {table_count} |
| Total unique codes | {len(all_records)} |
| 8-digit tariff lines | {len(lines8)} |
| 6-digit subheadings | {len(lines6)} |
| 4-digit headings | {len(lines4)} |
| Chapters with data | {len(bab_counts)} |

## Coverage per Bab
| Bab | Count |
|-----|-------|
"""
    for bab in sorted(bab_counts.keys()):
        report += f'| {bab} | {bab_counts[bab]} |\n'

    report += '\n## Key Code Verification\n'
    for code in test:
        status = 'OK' if code in by_code else 'MISSING'
        uid = by_code[code]['uraian_id'][:60] if code in by_code else '-'
        report += f'- {status}: `{code}` — {uid}\n'

    report += f'\n## Missing Bab\n{missing_bab}\n'

    with open(OUT_REPORT, 'w', encoding='utf-8') as f:
        f.write(report)
    print(f'\nReport: {OUT_REPORT}')
    print('Done.')

if __name__ == '__main__':
    main()
